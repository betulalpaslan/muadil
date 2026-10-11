"""Puan önerisi Excel'ini içe aktarıcının kurallarına göre denetler
(backend/Muadil.Infrastructure/IceAktarma/ExcelOkuyucu.cs ve IceAktarmaServisi.cs).

Kullanım:
  py -3.13 dogrula.py --excel D:/muadil/<dosya>.xlsx [--parfumler D:/muadil/muadil-48-gecici-puan.xlsx]

--parfumler verilirse orijinal marka + parfüm adlarının o dosyanın "Parfümler" sayfasında olup olmadığına da bakar
(toplu yükleme, adı veritabanında ya da yüklenen dosyanın Parfümler sayfasında olmayan satırı reddeder).
"""
import argparse
import sys

import openpyxl

ap = argparse.ArgumentParser()
ap.add_argument("--excel", required=True)
ap.add_argument("--parfumler")
a = ap.parse_args()

parfumler = None
if a.parfumler:
    p = openpyxl.load_workbook(a.parfumler, read_only=True)["Parfümler"]
    parfumler = {(r[0], r[1]) for r in p.iter_rows(min_row=2, values_only=True) if r[0]}

CINSIYETLER = {"", "kadın", "kadin", "erkek", "unisex"}

wb = openpyxl.load_workbook(a.excel)
hatalar, gorulen = [], set()
for no, r in enumerate(wb["Muadiller"].iter_rows(min_row=2, max_col=9, values_only=True), start=2):
    if not any(v not in (None, "") for v in r):
        continue
    om, op, mm, kod, fiyat, kal, ben, link, cinsiyet = r
    if not om or not op or not mm:
        hatalar.append((no, "marka/parfüm/muadil marka boş"))
    if not kod or len(str(kod)) > 50:
        hatalar.append((no, "kod boş ya da 50 karakterden uzun"))
    if not isinstance(fiyat, (int, float)) or not 1 <= fiyat <= 100000:
        hatalar.append((no, f"fiyat {fiyat}"))
    for ad, v in (("kalıcılık", kal), ("benzerlik", ben)):
        if not isinstance(v, int) or not 1 <= v <= 10:
            hatalar.append((no, f"{ad} {v} (1–10 tam sayı olmalı)"))
    if not link or not str(link).startswith(("http://", "https://")) or len(str(link)) > 500:
        hatalar.append((no, "link"))
    if str(cinsiyet or "").strip().lower() not in CINSIYETLER:
        hatalar.append((no, f"cinsiyet '{cinsiyet}' (Kadın/Erkek/Unisex/boş olmalı)"))
    anahtar = (str(mm).lower(), str(kod).lower(), str(cinsiyet or "unisex").strip().lower())
    if anahtar in gorulen:
        hatalar.append((no, f"{mm} {kod} ({cinsiyet or 'Unisex'}) dosyada birden fazla"))
    gorulen.add(anahtar)
    if parfumler is not None and (om, op) not in parfumler:
        hatalar.append((no, f"'{om} {op}' Parfümler listesinde yok"))

print(f"Muadiller: {len(gorulen)} satır, {len(hatalar)} hata")
for h in hatalar:
    print("  satır", *h)
sys.exit(1 if hatalar else 0)
